from io import BytesIO

import fitz
from docx import Document
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from backend.document_parser import extract_embedded_images


def _png_bytes() -> bytes:
    image = Image.new("RGB", (120, 80), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((10, 10, 110, 70), outline="black")
    draw.text((20, 30), "CloudRAG", fill="black")
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def test_extract_embedded_images_from_pdf():
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_image(fitz.Rect(50, 50, 250, 170), stream=_png_bytes())
    content = pdf.tobytes()
    pdf.close()
    assert len(extract_embedded_images("sample.pdf", content)) == 1


def test_extract_embedded_images_from_docx():
    document = Document()
    document.add_picture(BytesIO(_png_bytes()))
    buffer = BytesIO()
    document.save(buffer)
    assert len(extract_embedded_images("sample.docx", buffer.getvalue())) == 1


def test_extract_embedded_images_from_pptx():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(BytesIO(_png_bytes()), Inches(1), Inches(1))
    buffer = BytesIO()
    presentation.save(buffer)
    assert len(extract_embedded_images("sample.pptx", buffer.getvalue())) == 1
