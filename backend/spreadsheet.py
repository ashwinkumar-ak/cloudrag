import csv
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook


def json_safe(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return str(value)


def format_excel_value(value, number_format: str | None = None) -> str:
    if value is None:
        return ""

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()

    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"

    if isinstance(value, (int, float, Decimal)):
        fmt = number_format or "General"
        fmt_upper = fmt.upper()

        if "$" in fmt or "[$$" in fmt:
            decimals = _decimal_places(fmt)
            return f"${float(value):,.{decimals}f}"

        if "%" in fmt:
            decimals = _decimal_places(fmt)
            return f"{float(value) * 100:.{decimals}f}%"

        if any(symbol in fmt_upper for symbol in ("DD", "MM", "YY")) and isinstance(value, (datetime, date)):
            return str(value)

        if isinstance(value, float):
            if value.is_integer():
                return str(int(value))
            return f"{value:,.10f}".rstrip("0").rstrip(".")

        return str(value)

    return str(value)


def _decimal_places(number_format: str) -> int:
    match = re.search(r"[.,](0+)", number_format)
    return len(match.group(1)) if match else 0


def normalize_header(value, index: int) -> str:
    text = str(value).strip() if value is not None else ""
    return text or f"Column {index}"


def parse_spreadsheet(filename: str, content: bytes) -> list[dict]:
    extension = Path(filename).suffix.lower()
    if extension in {".xlsx", ".xlsm"}:
        return _parse_excel(content)
    if extension == ".csv":
        return _parse_csv(content)
    raise ValueError(f"Unsupported spreadsheet type: {extension or 'unknown'}")


def _parse_excel(content: bytes) -> list[dict]:
    workbook = load_workbook(
        BytesIO(content),
        read_only=False,
        data_only=True,
    )
    rows = []

    try:
        for worksheet in workbook.worksheets:
            values = list(worksheet.iter_rows(values_only=False))
            if not values:
                continue

            header_row_index = _find_header_row(values)
            headers = [
                normalize_header(cell.value, index + 1)
                for index, cell in enumerate(values[header_row_index])
            ]

            for excel_row_number, row_cells in enumerate(
                values[header_row_index + 1:],
                start=header_row_index + 2,
            ):
                raw = {}
                display = {}
                has_value = False

                for index, header in enumerate(headers):
                    cell = row_cells[index] if index < len(row_cells) else None
                    value = cell.value if cell is not None else None
                    if value is not None and str(value).strip() != "":
                        has_value = True
                    raw[header] = json_safe(value)
                    display[header] = format_excel_value(
                        value,
                        cell.number_format if cell is not None else None,
                    )

                if has_value:
                    rows.append({
                        "sheet_name": worksheet.title,
                        "row_number": excel_row_number,
                        "data": raw,
                        "display_data": display,
                    })
    finally:
        workbook.close()

    return rows


def _find_header_row(values) -> int:
    for index, row in enumerate(values[:10]):
        non_empty = sum(
            1 for cell in row
            if cell.value is not None and str(cell.value).strip() != ""
        )
        if non_empty >= 2:
            return index
    return 0


def _parse_csv(content: bytes) -> list[dict]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("CSV file must be UTF-8 encoded.") from exc

    reader = csv.reader(StringIO(text))
    rows = list(reader)
    if not rows:
        return []

    headers = [normalize_header(value, i + 1) for i, value in enumerate(rows[0])]
    result = []

    for row_number, values in enumerate(rows[1:], start=2):
        padded = values + [""] * (len(headers) - len(values))
        data = {header: _coerce_csv_value(value) for header, value in zip(headers, padded)}
        display = {header: value.strip() for header, value in zip(headers, padded)}
        if any(str(value).strip() for value in values):
            result.append({
                "sheet_name": "CSV",
                "row_number": row_number,
                "data": data,
                "display_data": display,
            })

    return result


def _coerce_csv_value(value: str):
    text = value.strip()
    if not text:
        return ""
    try:
        if re.fullmatch(r"[-+]?\d+", text):
            return int(text)
        if re.fullmatch(r"[-+]?(?:\d+\.\d*|\d*\.\d+)", text):
            return float(text)
    except ValueError:
        pass
    return text


def normalize_column_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def find_column(data: dict, requested: str) -> str | None:
    target = normalize_column_name(requested)
    exact = {normalize_column_name(key): key for key in data}
    if target in exact:
        return exact[target]

    for normalized, original in exact.items():
        if target and (target in normalized or normalized in target):
            return original
    return None


def parse_literal(value: str):
    value = value.strip().strip('"\'').rstrip("?.!")
    if re.fullmatch(r"[-+]?\d+", value):
        return int(value)
    if re.fullmatch(r"[-+]?(?:\d+\.\d*|\d*\.\d+)", value):
        return float(value)
    return value


def values_equal(actual, expected) -> bool:
    if actual is None:
        return str(expected).strip() == ""
    if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
        return float(actual) == float(expected)
    return str(actual).strip().lower() == str(expected).strip().lower()


def numeric_value(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value))
        except InvalidOperation:
            return None
    text = str(value).replace(",", "").replace("$", "").strip()
    if text.endswith("%"):
        text = text[:-1]
    try:
        return Decimal(text)
    except InvalidOperation:
        return None
