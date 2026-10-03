from pathlib import Path
import zipfile
from io import BytesIO

import fitz
from docx import Document as WordDocument
from openpyxl import load_workbook
from pptx import Presentation


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".pdf",
    ".docx",
    ".pptx",
    ".xlsx",
    ".xlsm",
    ".csv",
    ".jpg",
    ".jpeg",
    ".png",
    ".webp",
}


class EmbeddedImage:
    def __init__(self, content: bytes, mime_type: str, label: str):
        self.content = content
        self.mime_type = mime_type
        self.label = label


def extract_embedded_images(filename: str, content: bytes) -> list[EmbeddedImage]:
    """Extract embedded raster images from PDF, DOCX, and PPTX files.

    Images are returned in document order where the source format exposes
    stable ordering. Duplicate PDF image xrefs are emitted only once.
    """
    extension = Path(filename).suffix.lower()
    if extension == '.pdf':
        return _extract_pdf_images(content)
    if extension == '.docx':
        return _extract_docx_images(content)
    if extension == '.pptx':
        return _extract_pptx_images(content)
    return []


def _mime_from_extension(extension: str) -> str | None:
    return {
        '.jpg': 'image/jpeg',
        '.jpeg': 'image/jpeg',
        '.png': 'image/png',
        '.webp': 'image/webp',
    }.get(extension.lower())


def _extract_pdf_images(content: bytes) -> list[EmbeddedImage]:
    images: list[EmbeddedImage] = []
    seen_xrefs: set[int] = set()
    try:
        document = fitz.open(stream=content, filetype='pdf')
        for page_number, page in enumerate(document, start=1):
            for image_index, image_info in enumerate(page.get_images(full=True), start=1):
                xref = image_info[0]
                if xref in seen_xrefs:
                    continue
                seen_xrefs.add(xref)
                extracted = document.extract_image(xref)
                image_bytes = extracted.get('image')
                extension = extracted.get('ext', '')
                mime_type = _mime_from_extension('.' + extension)
                if not image_bytes or not mime_type:
                    continue
                images.append(EmbeddedImage(
                    content=image_bytes,
                    mime_type=mime_type,
                    label=f'PDF page {page_number}, image {image_index}',
                ))
        document.close()
        return images
    except Exception as exc:
        raise DocumentParseError('Failed to extract images from PDF document.') from exc


def _extract_docx_images(content: bytes) -> list[EmbeddedImage]:
    images: list[EmbeddedImage] = []
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            media_names = [
                name for name in archive.namelist()
                if name.startswith('word/media/') and not name.endswith('/')
            ]
            for index, name in enumerate(sorted(media_names), start=1):
                extension = Path(name).suffix.lower()
                mime_type = _mime_from_extension(extension)
                if not mime_type:
                    continue
                images.append(EmbeddedImage(
                    content=archive.read(name),
                    mime_type=mime_type,
                    label=f'DOCX embedded image {index}',
                ))
        return images
    except Exception as exc:
        raise DocumentParseError('Failed to extract images from Word document.') from exc


def _extract_pptx_images(content: bytes) -> list[EmbeddedImage]:
    images: list[EmbeddedImage] = []
    try:
        presentation = Presentation(BytesIO(content))
        seen: set[str] = set()
        for slide_number, slide in enumerate(presentation.slides, start=1):
            for image_index, shape in enumerate(slide.shapes, start=1):
                if not getattr(shape, 'shape_type', None) or not hasattr(shape, 'image'):
                    continue
                image = shape.image
                blob = image.blob
                extension = '.' + image.ext
                mime_type = _mime_from_extension(extension)
                if not mime_type:
                    continue
                key = image.sha1
                if key in seen:
                    continue
                seen.add(key)
                images.append(EmbeddedImage(
                    content=blob,
                    mime_type=mime_type,
                    label=f'PowerPoint slide {slide_number}, image {image_index}',
                ))
        return images
    except Exception as exc:
        raise DocumentParseError('Failed to extract images from PowerPoint presentation.') from exc


class DocumentParseError(RuntimeError):
    """Raised when a document cannot be parsed."""


def extract_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError(
            f"Unsupported file type: {extension or 'unknown'}"
        )

    if extension in {".txt", ".md"}:
        return _extract_text_file(content)

    if extension == ".pdf":
        return _extract_pdf(content)

    if extension == ".docx":
        return _extract_docx(content)

    if extension == ".pptx":
        return _extract_pptx(content)

    if extension in {".xlsx", ".xlsm"}:
        return _extract_excel(content)

    if extension == ".csv":
        return _extract_csv(content)

    if extension in {".jpg", ".jpeg", ".png", ".webp"}:
        return "[Image document]"

    raise DocumentParseError(
        f"Unsupported file type: {extension}"
    )


def _extract_text_file(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise DocumentParseError(
            "Text file must be UTF-8 encoded."
        ) from exc


def _extract_pdf(content: bytes) -> str:
    try:
        document = fitz.open(
            stream=content,
            filetype="pdf",
        )

        pages = []

        for page_number, page in enumerate(document, start=1):
            text = page.get_text().strip()

            if text:
                pages.append(
                    f"[Page {page_number}]\n{text}"
                )

        document.close()

        return "\n\n".join(pages)

    except Exception as exc:
        raise DocumentParseError(
            "Failed to read PDF document."
        ) from exc


def _extract_docx(content: bytes) -> str:
    try:
        from io import BytesIO

        document = WordDocument(
            BytesIO(content)
        )

        sections = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                sections.append(text)

        for table_index, table in enumerate(
            document.tables,
            start=1,
        ):
            sections.append(
                f"[Table {table_index}]"
            )

            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                sections.append(
                    " | ".join(cells)
                )

        return "\n".join(sections)

    except Exception as exc:
        raise DocumentParseError(
            "Failed to read Word document."
        ) from exc


def _extract_pptx(content: bytes) -> str:
    try:
        from io import BytesIO

        presentation = Presentation(
            BytesIO(content)
        )

        slides = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            slide_text = []

            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()

                    if text:
                        slide_text.append(text)

            if slide_text:
                slides.append(
                    f"[Slide {slide_number}]\n"
                    + "\n".join(slide_text)
                )

        return "\n\n".join(slides)

    except Exception as exc:
        raise DocumentParseError(
            "Failed to read PowerPoint presentation."
        ) from exc


def _extract_excel(content: bytes) -> str:
    try:
        from io import BytesIO

        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
        )

        sheets = []

        for worksheet in workbook.worksheets:
            rows = [
                f"[Sheet: {worksheet.title}]"
            ]

            for row in worksheet.iter_rows(
                values_only=True
            ):
                values = []

                for value in row:
                    if value is None:
                        values.append("")
                    else:
                        values.append(str(value))

                if any(value.strip() for value in values):
                    rows.append(
                        " | ".join(values)
                    )

            if len(rows) > 1:
                sheets.append(
                    "\n".join(rows)
                )

        workbook.close()

        return "\n\n".join(sheets)

    except Exception as exc:
        raise DocumentParseError(
            "Failed to read Excel workbook."
        ) from exc

def _extract_csv(content: bytes) -> str:
    try:
        from backend.spreadsheet import parse_spreadsheet

        rows = parse_spreadsheet("data.csv", content)
        sections = []

        for row in rows:
            values = [
                f"{key}: {value}"
                for key, value in row["display_data"].items()
                if value != ""
            ]
            sections.append(
                f"[Sheet: CSV, Row {row['row_number']}]\n"
                + " | ".join(values)
            )

        return "\n\n".join(sections)
    except Exception as exc:
        raise DocumentParseError(
            "Failed to read CSV document."
        ) from exc
